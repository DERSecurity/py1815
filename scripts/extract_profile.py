"""Read the IEEE 1815.2-2025 Profile Companion Data Point Tables into JSON.

IEEE Std 1815.2-2025 specifies its DNP3 points in a spreadsheet, the Profile
Companion Data Point Tables (CDPT), which the standard declares normative. IEEE
distributes it without charge alongside the standard:

    https://standards.ieee.org/wp-content/uploads/import/download/1815.2-2025_downloads.zip

This script turns that workbook into `conformance/ieee-1815-2-2025.json`, the
machine-readable form the profile machinery reads:

    python scripts/extract_profile.py --cdpt "<the workbook>.xlsx" --write
    python scripts/extract_profile.py --cdpt "<the workbook>.xlsx" --check

What is kept is what an implementation needs: every point's index (absolute, or
relative to a named block), name, event class, range, scaling, units, the IEC
61850 attribute it originates from, its paired point, purpose and mandatory
flags, plus the block-start table from the Key sheet. What is dropped is the
per-device "EUT capabilities" columns, which are a blank template for a device
profile document and carry nothing about the profile itself.

Columns are located by header text and checked, so a workbook whose layout has
changed fails here rather than mis-mapping a column. `openpyxl` is not a
dependency of this package; install it just to run this script.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import pathlib
import re
import sys
from typing import Any

#: Where the checked-in table lives, relative to the repository root.
TABLES = pathlib.Path(__file__).resolve().parent.parent / "conformance" / "ieee-1815-2-2025.json"

EDITION = "IEEE Std 1815.2-2025"
SOURCE_URL = (
    "https://standards.ieee.org/wp-content/uploads/import/download/1815.2-2025_downloads.zip"
)

KINDS = ("BO", "BI", "AO", "AI", "CTR")

#: An absolute index cell: `AI328`, `CTR0`.
_ABSOLUTE = re.compile(r"^(?P<kind>AI|AO|BI|BO|CTR)\s*(?P<index>\d+)$")

#: A block-relative index cell: `Meter_AI+3`, `Meter_BI + 1`, `Exp_BO+10`, or
#: bare `Inverter_BI` for offset zero.
_RELATIVE = re.compile(
    r"^(?P<block>[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)*?)_(?P<kind>AI|AO|BI|BO|CTR)"
    r"(?:\s*\+\s*(?P<offset>\d+))?$"
)

#: What a section heading's first cell may hold when it is not empty: a block's
#: start index (`5000`), a bare block symbol (`Exp`, `VP`), a heading in words
#: (`System Counters`, `Not currently used`), or the `. . .` that stands for
#: the units between the first and the last. No digit beside a letter, no `+`,
#: no `_`, no parenthesis: those belong to an index, and an index that fails
#: to parse is an error rather than a heading.
_HEADING = re.compile(r"^(?:\d+|[A-Za-z][A-Za-z ]*|\.(?:\s*\.)*)$")

#: The equipment sheets state each per-unit block twice: once for unit #1
#: (`Meter_AI+3`) and once for a generic unit #m (`Meter_AI+mhai(m)+3`), where
#: the function names the unit's block start. The second is the first restated
#: with an instance term, so it is checked against the first and not kept. Its
#: spelling drifts (`Battery_bhai(b)+1` has no kind; `HM_CTR` is an alias), so
#: the kind is optional here and taken from the sheet.
_GENERIC = re.compile(
    r"^(?P<block>[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)*?)_(?:(?P<kind>AI|AO|BI|BO|CTR))?"
    r"\s*\+?\s*[A-Za-z]+\([a-z]\)\s*(?:\+\s*(?P<offset>\d+))?$"
)

#: The counter sheet names the meter block `HM`, AN2018-001's symbol for the
#: historical meter block, where every other sheet says `Meter`; one paired
#: reference writes the experimental block as `EX`. One symbol per block, so
#: the loader has one thing to resolve.
_BLOCK_ALIASES = {"HM": "Meter", "EX": "Exp"}

#: The paired-point columns write some equipment references kind-first:
#: `AO+Meter+3` where the index column would say `Meter_AO+3`. Same meaning,
#: second spelling.
_REFERENCE_KIND_FIRST = re.compile(
    r"^(?P<kind>AI|AO|BI|BO|CTR)\s*\+\s*(?P<block>[A-Za-z][A-Za-z0-9_]*?)\s*\+\s*(?P<offset>\d+)$"
)

#: Header text expected in each sheet, in column order, as substrings of the
#: two header rows joined. The extractor refuses a sheet whose headers do not
#: match, which is how a reordered edition fails loudly.
_HEADERS: dict[str, list[tuple[str, str]]] = {
    "BO": [
        ("index", "Point Index"),
        ("name", "Name / Description"),
        ("state_0", "value is 0"),
        ("state_1", "value is 1"),
        ("iec61850", "IEC61850UniqueString"),
        ("associated", "Assoc. BI"),
        ("purpose", "Purpose"),
        ("mandatory_1815_2", "Mandatory for IEEE 1815.2"),
        ("mandatory_1547", "Mandatory for 1547"),
    ],
    "BI": [
        ("index", "Point Index"),
        ("name", "Name / Description"),
        ("event_class", "Default Event Class"),
        ("state_0", "value is 0"),
        ("state_1", "value is 1"),
        ("iec61850", "IEC61850UniqueString"),
        ("associated", "Assoc. BO"),
        ("purpose", "Purpose"),
        ("mandatory_1815_2", "Mandatory for IEEE 1815.2"),
        ("mandatory_1547", "Mandatory for 1547"),
    ],
    "AO": [
        ("index", "Point Index"),
        ("name", "Name / Description"),
        ("minimum", "Minimum"),
        ("maximum", "Maximum"),
        ("multiplier", "Multiplier"),
        ("offset", "OffSet"),
        ("units", "Units"),
        ("iec61850", "IEC61850UniqueString"),
        ("associated", "Assoc. AI"),
        ("purpose", "Purpose"),
        ("mandatory_1815_2", "Mandatory for IEEE 1815.2"),
        ("mandatory_1547", "Mandatory for 1547"),
    ],
    "AI": [
        ("index", "Point Index"),
        ("name", "Name / Description"),
        ("event_class", "Default Event Class"),
        ("minimum", "Minimum"),
        ("maximum", "Maximum"),
        ("multiplier", "Multiplier"),
        ("offset", "Offset"),
        ("units", "Units"),
        ("iec61850", "IEC61850UniqueString"),
        ("value", "Value"),
        ("associated", "Assoc. AO"),
        ("purpose", "Purpose"),
        ("mandatory_1815_2", "Mandatory for IEEE 1815.2"),
        ("mandatory_1547", "Mandatory for 1547"),
    ],
    "CTR": [
        ("index", "Point Index"),
        ("name", "Name / Description"),
        ("event_class", "Default Class Assigned to"),
        ("frozen", "Frozen Counter Exists"),
        ("frozen_event_class", "Default Class Assigned to"),
        ("iec61850", "IEC61850UniqueString"),
        ("purpose", "Purpose"),
        ("mandatory_1815_2", "Mandatory for IEEE 1815.2"),
        ("mandatory_1547", "Mandatory for 1547"),
    ],
}


class ExtractionError(RuntimeError):
    """The workbook did not yield what this script expects to find."""


def _text(cell: Any) -> str:
    return "" if cell is None else " ".join(str(cell).split())


def _number(cell: Any) -> int | float | None:
    """A numeric cell as a number, an empty or textual cell as None."""
    if cell is None or cell == "":
        return None
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return int(cell) if float(cell).is_integer() else cell
    text = _text(cell).replace(",", "")
    try:
        value = float(text)
    except ValueError:
        return None
    return int(value) if value.is_integer() else value


def _event_class(cell: Any) -> tuple[int | None, str | None]:
    """A class number, or None and the note the cell carried instead."""
    number = _number(cell)
    if number is not None and number in (0, 1, 2, 3):
        return int(number), None
    text = _text(cell)
    return None, (text or None)


def _flag(cell: Any) -> bool:
    return _text(cell).upper().startswith("M")


def _optional(cell: Any) -> str | None:
    text = _text(cell)
    return text or None


def _check_headers(sheet: Any, kind: str) -> None:
    rows = list(sheet.iter_rows(min_row=1, max_row=2, values_only=True))
    if len(rows) < 2:
        raise ExtractionError(f"{kind}: fewer than two header rows")
    for column, (field, expected) in enumerate(_HEADERS[kind]):
        cells = [_text(rows[0][column]) if column < len(rows[0]) else ""]
        cells.append(_text(rows[1][column]) if column < len(rows[1]) else "")
        joined = " ".join(cells)
        if expected.lower() not in joined.lower():
            raise ExtractionError(
                f"{kind}: column {column + 1} should carry {field!r} ({expected!r}); "
                f"header reads {joined!r}"
            )


def _parse_index(cell: Any, kind: str) -> dict[str, Any] | None:
    """The index a cell names, or None when the cell is not a point index."""
    text = _text(cell)
    if not text:
        return None
    absolute = _ABSOLUTE.match(text)
    if absolute:
        if absolute.group("kind") != kind:
            raise ExtractionError(f"{kind}: index {text!r} names another kind")
        return {"index": int(absolute.group("index")), "block": None, "offset": None}
    relative = _RELATIVE.match(text)
    if relative:
        if relative.group("kind") != kind:
            raise ExtractionError(f"{kind}: index {text!r} names another kind")
        block = relative.group("block")
        return {
            "index": None,
            "block": _BLOCK_ALIASES.get(block, block),
            "offset": int(relative.group("offset") or 0),
        }
    return None


def _reference(cell: Any) -> str | None:
    """A paired-point cell in the spelling the index column uses.

    The workbook writes the same reference two ways -- `Meter_AO+3` in the
    index column and `AO+Meter+3` in the paired-point column -- and once
    writes a block by an alias. One spelling means a loader resolves one
    grammar. A cell that matches neither is kept as written rather than
    dropped, so nothing the workbook says is lost, only unnormalized.
    """
    text = _text(cell)
    if not text:
        return None
    absolute = _ABSOLUTE.match(text)
    if absolute:
        return f"{absolute.group('kind')}{int(absolute.group('index'))}"
    relative = _RELATIVE.match(text) or _REFERENCE_KIND_FIRST.match(text)
    if relative:
        block = relative.group("block")
        block = _BLOCK_ALIASES.get(block, block)
        return f"{block}_{relative.group('kind')}+{int(relative.group('offset') or 0)}"
    return text


def _points(workbook: Any, kind: str) -> list[dict[str, Any]]:
    """Every point row of one sheet, with the section heading it sits under."""
    sheet = workbook[kind]
    _check_headers(sheet, kind)
    fields = [field for field, _ in _HEADERS[kind]]
    width = len(fields)
    points: list[dict[str, Any]] = []
    #: Offsets the generic-unit restatement of each block names, to be checked
    #: against the unit #1 rows once the sheet is read.
    restated: dict[str, set[int]] = {}
    section = None
    for row in sheet.iter_rows(min_row=3, max_col=width, values_only=True):
        cells = list(row) + [None] * (width - len(row))
        named = dict(zip(fields, cells, strict=True))
        where = _parse_index(named["index"], kind)
        if where is None:
            # A row that is not a point is one of exactly three things: the
            # generic-unit restatement of a block, a heading or note, or a
            # blank line. Anything else in the index cell is an index the
            # grammar did not recognize -- a typo, or a later edition's
            # spelling -- and dropping it would lose a point without a trace.
            first = _text(named["index"])
            generic = _GENERIC.match(first)
            if generic:
                block = _BLOCK_ALIASES.get(generic.group("block"), generic.group("block"))
                restated.setdefault(block, set()).add(int(generic.group("offset") or 0))
                continue
            if first and not _HEADING.match(first):
                raise ExtractionError(f"{kind}: index cell {first!r} is not a point or a heading")
            heading = _text(named["name"])
            if heading:
                section = heading
            continue
        point: dict[str, Any] = {
            **where,
            "name": _text(named["name"]),
            "section": section,
        }
        if "event_class" in named:
            point["event_class"], note = _event_class(named["event_class"])
            if note:
                point["event_class_note"] = note
        if "state_0" in named:
            point["states"] = [_text(named["state_0"]), _text(named["state_1"])]
        if "minimum" in named:
            point["minimum"] = _number(named["minimum"])
            point["maximum"] = _number(named["maximum"])
            point["multiplier"] = _number(named["multiplier"])
            point["offset_value"] = _number(named["offset"])
            point["units"] = _optional(named["units"])
        if "value" in named:
            point["value"] = _optional(named["value"])
        if "frozen" in named:
            point["frozen"] = _text(named["frozen"]).lower().startswith("y")
            point["frozen_event_class"], _ = _event_class(named["frozen_event_class"])
        point["iec61850"] = _optional(named["iec61850"])
        if "associated" in named:
            point["associated"] = _reference(named["associated"])
        point["purpose"] = _optional(named["purpose"])
        point["mandatory_1815_2"] = _flag(named["mandatory_1815_2"])
        point["mandatory_1547"] = _flag(named["mandatory_1547"])
        points.append(point)
    if not points:
        raise ExtractionError(f"{kind}: no point rows found")
    # The two statements of each block have to agree, or the sheet is telling
    # two stories about how long the block is; the unit #1 rows are the ones
    # kept, so a disagreement is an error rather than a choice.
    for block, offsets in restated.items():
        first_unit = {p["offset"] for p in points if p["block"] == block}
        if offsets != first_unit:
            raise ExtractionError(
                f"{kind}: block {block} is restated for a generic unit with offsets "
                f"{sorted(offsets ^ first_unit)} that its unit #1 rows do not match"
            )
    return points


def _key(workbook: Any) -> dict[str, Any]:
    """The Key sheet's block-start table.

    The sheet is laid out for reading, not parsing: the table sits in columns
    I through P, block names in the first of those, then a start index and one
    column per point kind. Rows are matched by their label.
    """
    sheet = workbook["Key"]
    rows = [list(r) for r in sheet.iter_rows(min_col=9, max_col=16, values_only=True)]

    def find(label: str) -> list[Any]:
        for row in rows:
            if _text(row[0]).lower().startswith(label.lower()):
                return row
        raise ExtractionError(f"Key: no row labelled {label!r}")

    def starts(row: list[Any]) -> dict[str, int]:
        # Columns: label, note, start, BO, BI, AO, AI, CTR.
        out = {}
        for kind, cell in zip(KINDS, row[3:8], strict=True):
            number = _number(cell)
            if number is not None:
                out[kind] = int(number)
        return out

    fixed = {}
    for label, key in [
        ("Configuration & Control", "configuration"),
        ("Functions", "functions"),
        ("Curves", "curves"),
        ("System Meter", "system_meter"),
        ("Extensions beyond", "extensions"),
        ("Gap #1", "gap_1"),
        ("Schedules (2 items", "schedules_2"),
        ("Schedule Status (2", "schedule_status_2"),
        ("Schedules (4 items", "schedules_4"),
        ("Schedule Status (4", "schedule_status_4"),
    ]:
        fixed[key] = starts(find(label))

    equipment = {}
    for label, key, symbol in [
        ("Meters", "meters", "Meter"),
        ("DER Units", "der_units", "DER_Unit"),
        ("Inverters", "inverters", "Inverter"),
        ("Batteries", "batteries", "Battery"),
    ]:
        row = find(label)
        start = _number(row[2])
        if start is None:
            raise ExtractionError(f"Key: {label} has no start index")
        # The row after the block's own is its per-unit block length.
        position = rows.index(row)
        per_unit = starts(rows[position + 1])
        equipment[key] = {"symbol": symbol, "start": int(start), "per_unit": per_unit}

    def single(label: str) -> int:
        number = _number(find(label)[2])
        if number is None:
            raise ExtractionError(f"Key: {label} has no start index")
        return int(number)

    return {
        "fixed": fixed,
        "equipment": equipment,
        "experimental": {"symbol": "Exp", "start": single("Experimental")},
        "vendor": {"symbol": "VP", "start": single("Vendor Points")},
        "start_index_points": single("Start Index Points"),
        "maximum_index": single("Maximum Index"),
    }


def extract(path: pathlib.Path) -> dict[str, Any]:
    """The whole workbook as one document."""
    try:
        import openpyxl
    except ModuleNotFoundError:  # pragma: no cover - depends on the caller's venv
        raise SystemExit("this script needs openpyxl: pip install openpyxl") from None
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    missing = [name for name in ("Key", *KINDS) if name not in workbook.sheetnames]
    if missing:
        raise ExtractionError(f"workbook lacks sheets {missing}; is this the CDPT?")

    key = _key(workbook)
    points = {kind: _points(workbook, kind) for kind in KINDS}

    # The Key sheet states each equipment block's per-unit length, and three
    # of the analog input blocks hold one more point than it says. The rows
    # are the profile, so the length the rows imply is recorded beside the
    # Key's, and a loader sizing a block should use the larger.
    for entry in key["equipment"].values():
        observed = {}
        for kind in KINDS:
            offsets = [p["offset"] for p in points[kind] if p["block"] == entry["symbol"]]
            if offsets:
                observed[kind] = max(offsets) + 1
        entry["per_unit_observed"] = observed
    # The analog input sheet ends with a "Maximum Points" row at the highest
    # index the space allows. It marks the end of the space rather than naming
    # a point, so it is not one.
    points["AI"] = [p for p in points["AI"] if p["index"] != key["maximum_index"]]

    version = next(
        (p["value"] for p in points["AI"] if p["index"] == 0 and p.get("value")),
        None,
    )
    # Clause 5.2 says the block starting indices are published in band from
    # analog input 65000, and the point rows agree; the Key sheet's own figure
    # for that block does not. The rows are the profile, so what they say is
    # recorded here and the Key's figure stays where the Key put it.
    advertised = [
        p["index"]
        for p in points["AI"]
        if p["index"] is not None and p["section"] and "auto discovery" in p["section"].lower()
    ]
    if not advertised:
        raise ExtractionError("AI: no auto-discovery index points found")
    # Recorded as a first index and a count, which describes the block only
    # if it has no holes; a missing row would otherwise be papered over as a
    # shorter block that still ends one short of the real last point.
    expected = list(range(min(advertised), min(advertised) + len(advertised)))
    if sorted(advertised) != expected:
        missing = sorted(set(expected) - set(advertised))
        raise ExtractionError(f"AI: auto-discovery block is not contiguous; missing {missing}")
    return {
        "edition": EDITION,
        "generated_by": "scripts/extract_profile.py",
        "source": {
            "title": "IEEE 1815.2 Profile Companion Data Point Tables",
            "url": SOURCE_URL,
            "file": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "profile_version": version,
        },
        "key": key,
        "index_points": {"first": min(advertised), "count": len(advertised)},
        "points": points,
    }


def _serialize(document: dict[str, Any]) -> str:
    return json.dumps(document, indent=1, ensure_ascii=False) + "\n"


def _summary(document: dict[str, Any]) -> str:
    counts = ", ".join(f"{kind} {len(document['points'][kind])}" for kind in KINDS)
    relative = sum(
        1 for kind in KINDS for point in document["points"][kind] if point["index"] is None
    )
    return (
        f"profile version {document['source']['profile_version']}; points: {counts}; "
        f"{relative} of them block-relative"
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--cdpt", required=True, type=pathlib.Path, help="the CDPT workbook")
    action = parser.add_mutually_exclusive_group()
    action.add_argument("--write", action="store_true", help=f"overwrite {TABLES.name}")
    action.add_argument(
        "--check", action="store_true", help="fail if the workbook and the JSON disagree"
    )
    args = parser.parse_args(argv)

    if not args.cdpt.is_file():
        print(f"no such file: {args.cdpt}", file=sys.stderr)
        return 2
    try:
        document = extract(args.cdpt)
    except ExtractionError as error:
        print(f"extraction failed: {error}", file=sys.stderr)
        return 1

    if args.write:
        TABLES.parent.mkdir(parents=True, exist_ok=True)
        TABLES.write_text(_serialize(document), encoding="utf-8")
        print(f"wrote {TABLES}: {_summary(document)}")
        return 0

    if args.check:
        if not TABLES.is_file():
            print(f"{TABLES} does not exist; run with --write", file=sys.stderr)
            return 1
        current = json.loads(TABLES.read_text(encoding="utf-8"))
        # The hash names the exact workbook; a re-download that differs only in
        # metadata would fail the check for no reason, so compare the content.
        current["source"].pop("sha256", None)
        fresh = dict(document, source=dict(document["source"]))
        fresh["source"].pop("sha256", None)
        if current == fresh:
            print(f"{TABLES.name} matches {args.cdpt.name}: {_summary(document)}")
            return 0
        print(f"{TABLES.name} does not match {args.cdpt.name}; run with --write", file=sys.stderr)
        return 1

    print(_summary(document))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
